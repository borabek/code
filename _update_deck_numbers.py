# -*- coding: utf-8 -*-
"""Slaytlardaki SAYILARI bugun olculenlerle esitle. YENI SLAYT YOK, yerlesim degismez.

Kaynak: results/p9_hizalama.json (100 parca, sizintisiz, gercek 4-model hatti) ve
cp_config.json > current_product (dagitilan ayarlar).

Bicimi korumak icin metin ILK RUN'a yazilir, kalan run'lar silinir -- shape.text'e atamak
yazi tipini/rengini sifirlardi.
"""
from pptx import Presentation
SRC = "ConnectionPointDetector_Bora_Bayrakci.pptx"


def set_text(sh, new):
    p = sh.text_frame.paragraphs[0]
    if not p.runs:
        return False
    p.runs[0].text = new
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    return True


def edit(prs, slide_no, old, new, near=None):
    """slide_no (1-tabanli) uzerinde metni 'old' olan sekli bulup 'new' yaz."""
    s = prs.slides[slide_no - 1]
    hits = [sh for sh in s.shapes
            if sh.has_text_frame and sh.text_frame.text.strip() == old
            and (near is None or abs(sh.top / 914400 - near) < 0.06)]
    if len(hits) != 1:
        print(f"  ATLANDI s{slide_no}: {old!r} -> {len(hits)} eslesme bulundu")
        return False
    set_text(hits[0], new)
    print(f"  s{slide_no}: {old!r} -> {new!r}")
    return True


def main():
    prs = Presentation(SRC)
    n = 0

    # --- CP-F1 basligi: 76.0% -> 80.2% (P9, 2026-07-30) ---
    n += edit(prs, 6, "76.0%", "80.2%")                      # tez karsilastirmasi
    n += edit(prs, 10, "76.0%", "80.2%", near=1.92)          # "The running system"

    # --- rejim tablosu (slayt 11) ---
    n += edit(prs, 11, "77.5%", "81.7%")
    n += edit(prs, 11, "62.6%", "67.2%")

    # --- gate artik tam korpusla egitildi (slayt 4) ---
    n += edit(prs, 4, "1 041 parts", "1 906 parts", near=3.88)

    # --- dagitilan ayarlar (slayt 8) ---
    n += edit(prs, 8, "10 vertices", "4 vertices")
    n += edit(prs, 8, "35%", "40%", near=2.74)               # dusuk-CP esigi
    n += edit(prs, 8, "25%", "35%", near=3.10)               # cok-CP esigi
    n += edit(prs, 8, "Auto-accept confidence", "Auto-accept gate score")
    n += edit(prs, 8, "50%", "60%", near=3.46)

    # --- sonraki adimlar: biri BITTI, biri OLCULUP OLDU (slayt 13) ---
    n += edit(prs, 13, "Retrain the filter on all 1 906 parts",
              "Retrain the filter on all 1 906 parts")
    n += edit(prs, 13, "it still runs on 1 041", "+12.7 points, measured")
    n += edit(prs, 13, "running now", "done")
    n += edit(prs, 13, "+2.8 points, code is ready",
              "+2.8 points on the older base; needs re-measuring")
    n += edit(prs, 13, "Higher mesh resolution for high-CP-density parts",
              "Higher mesh resolution for high-CP-density parts")
    n += edit(prs, 13, "recovers sub-threshold openings",
              "measured: high CP +8.1, low CP -10.8, net -8.8")
    n += edit(prs, 13, "designed", "dropped")

    prs.save(SRC)
    print(f"\n{n} alan guncellendi -> {SRC}")


if __name__ == "__main__":
    main()
