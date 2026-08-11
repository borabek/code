# -*- coding: utf-8 -*-
"""Slaytlari DISKTEN geri okuyup dogrula: cakisma, tasma, oran bozulmasi.

Cizen kod kendini dogrulayamaz -- ayni ilke CIZIM_SOZLESMESI'nde GLB icin de gecerli.
"""
import sys, io
from pptx import Presentation
from PIL import Image

FOOT_T = 7.00      # alt bant burada basliyor
TITLE_B = 1.30     # baslik burada bitiyor


def boxes(s):
    out = []
    for sh in s.shapes:
        if sh.top is None or sh.left is None:
            continue
        t = sh.top / 914400; l = sh.left / 914400
        w = sh.width / 914400; h = sh.height / 914400
        if t >= FOOT_T - 0.01:          # marka bandi / sayfa no -- kapsam disi
            continue
        txt = sh.text_frame.text.replace("\n", " ")[:34] if sh.has_text_frame else ""
        kind = "PIC" if sh.shape_type == 13 else "TXT"
        out.append((kind, l, t, w, h, txt, sh))
    return out


def overlap(a, b):
    _, l1, t1, w1, h1, *_ = a
    _, l2, t2, w2, h2, *_ = b
    ox = min(l1 + w1, l2 + w2) - max(l1, l2)
    oy = min(t1 + h1, t2 + h2) - max(t1, t2)
    return ox > 0.05 and oy > 0.05, round(max(ox, 0), 2), round(max(oy, 0), 2)


def main():
    prs = Presentation("ConnectionPointDetector_Bora_Bayrakci.pptx")
    bad = 0
    for idx in (11, 12):
        s = prs.slides[idx - 1]
        B = [b for b in boxes(s) if b[4] > 0.03]     # ayrac cizgilerini es gec
        print(f"--- SLAYT {idx} ---")
        for i in range(len(B)):
            for j in range(i + 1, len(B)):
                ov, ox, oy = overlap(B[i], B[j])
                if ov:
                    bad += 1
                    print(f"  CAKISMA {ox}x{oy}in: '{B[i][5]}' <-> '{B[j][5]}'")
        for k, l, t, w, h, txt, _sh in B:
            if t + h > FOOT_T + 0.01:
                bad += 1; print(f"  ALT BANDA TASIYOR ({t+h:.2f}in): '{txt}'")
            if t < TITLE_B - 0.01 and t > 0.9:
                bad += 1; print(f"  BASLIGA GIRIYOR ({t:.2f}in): '{txt}'")
            if l < 0.3 or l + w > 13.05:
                bad += 1; print(f"  KENARDAN TASIYOR (L{l:.2f} R{l+w:.2f}): '{txt}'")
        for sh in s.shapes:
            if sh.shape_type == 13:
                im = Image.open(io.BytesIO(sh.image.blob))
                nat = im.size[0] / im.size[1]; cur = sh.width / sh.height
                err = abs(cur - nat) / nat
                flag = "OK" if err < 0.02 else "ORAN BOZUK"
                if err >= 0.02: bad += 1
                print(f"  resim {sh.width/914400:.2f}x{sh.height/914400:.2f}in  "
                      f"oran {cur:.3f} / dogal {nat:.3f}  {flag}")
        print(f"  en alt icerik: {max(t+h for *_ , t, w, h, txt, sh in [(b[0],b[1],b[2],b[3],b[4],b[5],b[6]) for b in B]):.2f}in (alt bant {FOOT_T})")
    print(f"\n{'GECTI -- cakisma/tasma yok' if bad == 0 else f'{bad} SORUN VAR'}")
    return 1 if bad else 0


sys.exit(main())
